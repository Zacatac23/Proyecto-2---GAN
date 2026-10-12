"""Version Word editable con el mismo contenido y evidencia del informe PDF."""
import sys
from pathlib import Path
from html.parser import HTMLParser
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import build_report


def font(run,size=11,bold=False):
    run.font.name='Arial';run.font.size=Pt(size);run.font.bold=bold
    run.font.color.rgb=RGBColor(0,0,0)
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),'Arial')
    return run


class RichText(HTMLParser):
    def __init__(self,p,size=11):
        super().__init__(convert_charrefs=True);self.p=p;self.size=size;self.bold=0;self.link=None
    def handle_starttag(self,tag,attrs):
        if tag=='b':self.bold+=1
        if tag=='link':self.link=dict(attrs).get('href')
        if tag=='br':self.p.add_run().add_break()
    def handle_endtag(self,tag):
        if tag=='b':self.bold-=1
        if tag=='link':self.link=None
    def handle_data(self,text):
        if not self.link:font(self.p.add_run(text),self.size,self.bold>0);return
        rel=self.p.part.relate_to(self.link,RT.HYPERLINK,is_external=True)
        h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel)
        r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
        rf=OxmlElement('w:rFonts');rf.set(qn('w:ascii'),'Arial');rf.set(qn('w:hAnsi'),'Arial');pr.append(rf)
        sz=OxmlElement('w:sz');sz.set(qn('w:val'),str(int(self.size*2)));pr.append(sz)
        color=OxmlElement('w:color');color.set(qn('w:val'),'000000');pr.append(color)
        r.append(pr);t=OxmlElement('w:t');t.text=text;r.append(t);h.append(r);self.p._p.append(h)


class EditableReport:
    def __init__(self,path):
        self.path=path;path.parent.mkdir(parents=True,exist_ok=True)
        self.doc=Document();self.page=0;self.w=468
        section=self.doc.sections[0]
        section.page_width=Inches(8.5);section.page_height=Inches(11)
        section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Inches(1)
        section.footer_distance=Inches(.5)
        for name in ['Normal','Title','Heading 1','Caption']:
            s=self.doc.styles[name];s.font.name='Arial';s.font.size=Pt(11);s.font.color.rgb=RGBColor(0,0,0)
            s.font.bold=name in ['Title','Heading 1'];s.font.italic=False
            s.paragraph_format.space_before=Pt(0);s.paragraph_format.space_after=Pt(8)
            s.paragraph_format.line_spacing_rule=WD_LINE_SPACING.EXACTLY;s.paragraph_format.line_spacing=Pt(15)
            for border in list(s.element.iter(qn('w:pBdr'))):
                border.getparent().remove(border)
        normal=self.doc.styles['Normal'];normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=font(footer.add_run(),10)
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');r._r.addnext(field)
        self.doc.core_properties.title='Diseño generativo de personajes Informe de resultados'
        self.doc.core_properties.author='Luis Gonzalez - 23353; Jonathan Zacarias - 231104'
        self.doc.core_properties.subject='Experimentos controlados con GAN y galería regenerable'
    def new(self,title):
        self.page+=1
        p=self.doc.add_paragraph(style='Title' if self.page==1 else 'Heading 1')
        p.paragraph_format.page_break_before=self.page>1;p.paragraph_format.space_after=Pt(12)
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        font(p.add_run(title),11,True)
    def p(self,text,gap=8):
        p=self.doc.add_paragraph();p.paragraph_format.space_after=Pt(gap)
        p.paragraph_format.keep_together=True
        RichText(p).feed(text)
    def image(self,path,maxheight=320,width=None):
        img=Image.open(path);w,h=img.size;ratio=min((width or self.w)/w,maxheight/h)
        p=self.doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.SINGLE
        p.paragraph_format.space_after=Pt(9)
        run=p.add_run();font(run,1)
        run.add_picture(str(path),width=Pt(w*ratio),height=Pt(h*ratio))
    def spacer(self,height=10):
        p=self.doc.add_paragraph();p.paragraph_format.line_spacing=Pt(height)
        p.paragraph_format.space_after=Pt(0);font(p.add_run(''),1)
    def base_table(self,rows,widths,borders=False):
        t=self.doc.add_table(rows=rows,cols=len(widths));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
        for col,width in zip(t.columns,widths):col.width=Pt(width)
        pr=t._tbl.tblPr
        b=OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single' if borders else 'nil')
            if borders:e.set(qn('w:sz'),'4');e.set(qn('w:color'),'000000')
            b.append(e)
        pr.append(b)
        margins=OxmlElement('w:tblCellMar')
        for side in ['top','left','bottom','right']:
            e=OxmlElement('w:'+side);e.set(qn('w:w'),'120' if borders else '0');e.set(qn('w:type'),'dxa');margins.append(e)
        pr.append(margins)
        for row in t.rows:
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell,width in zip(row.cells,widths):
                cell.width=Pt(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0)
        return t
    def table(self,headers,rows,widths):
        t=self.base_table(len(rows)+1,widths,True)
        t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for i,row in enumerate([headers]+rows):
            for cell,text in zip(t.rows[i].cells,row):
                p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.line_spacing=Pt(13)
                font(p.add_run(str(text)),10.5,i==0)
        self.spacer(12)
    def cell_picture(self,cell,path,size):
        p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.SINGLE
        font(p.add_run(),1).add_picture(str(path),width=Pt(size),height=Pt(size))
    def caption(self,cell,text,bold=False,before=2,after=0):
        p=cell.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing=Pt(13);p.paragraph_format.space_before=Pt(before);p.paragraph_format.space_after=Pt(after)
        font(p.add_run(text),11 if bold else 10.5,bold)
    def sprites(self,paths,labels,columns=6,size=64):
        t=self.base_table((len(paths)+columns-1)//columns,[self.w/columns]*columns)
        for i,(path,label) in enumerate(zip(paths,labels)):
            cell=t.cell(i//columns,i%columns);self.cell_picture(cell,path,size);self.caption(cell,label,after=7)
        self.spacer(6)
    def gallery(self,directory,characters):
        t=self.base_table(2,[self.w/5]*5)
        for i,r in enumerate(characters):
            cell=t.cell(i//5,i%5);self.cell_picture(cell,directory/r['image'],78)
            self.caption(cell,r['name'],True,before=4)
            self.caption(cell,r['role']);self.caption(cell,f"Seed {r['seed']}",after=12)
        self.spacer(8)
    def neighbors(self,pairs):
        t=self.base_table(5,[117]*4)
        for i,(g,n,gl,nl) in enumerate(pairs):
            for j,(path,label) in enumerate([(g,gl),(n,nl)]):
                cell=t.cell(i//2,(i%2)*2+j);self.cell_picture(cell,path,72);self.caption(cell,label,before=3,after=15)
    def save(self,path):
        self.doc.save(path);assert self.page==12
        print('Word editable:',path)


if __name__=='__main__':
    build_report.Report=EditableReport
    if '--output' not in sys.argv:sys.argv+=['--output',str(build_report.ROOT/'presentacion_editable.docx')]
    build_report.main()
